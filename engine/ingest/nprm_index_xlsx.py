"""Read the challenge organizers' AD index (``nprm_index`` sheet).

This is the same 431 proposed rules as the FAA export that
:mod:`engine.ingest.nprm_xlsx` reads, re-published with a stable ``record_id``
per row and four derived columns. Everything that can be shared with that
module is imported from it: :class:`~engine.ingest.nprm_xlsx.NPRMRow`,
:func:`~engine.ingest.nprm_xlsx.clean_cell`,
:func:`~engine.ingest.nprm_xlsx.parse_date` and
:func:`~engine.ingest.nprm_xlsx.page_range`.

Which of the FAA export's verified quirks still hold in this file
---------------------------------------------------------------

STILL TRUE
  * 431 data rows, one per notice, 322 distinct dockets after normalization and
    99 of them re-noticed. One node per ROW, never per docket.
  * ``docket_number`` is still not uniformly FAA-YYYY-NNNN and two cells still
    name two dockets each (rows 374 and 383 of the sheet), so
    :func:`normalize_sheet_docket` is still required.
  * ``make`` is still unnormalized ("Boeing" vs "The Boeing Company"), so both
    the normalized and the raw string are carried out of here.

CHANGED BY THE ORGANIZERS
  * Dates are ISO ``%Y-%m-%d`` text, not ``%m/%d/%Y``. 431/431 parse.
  * The header is on row index 2: row 0 is a banner sentence and row 1 is
    blank. The FAA export has its header on row 0.
  * Columns are renamed and reordered (``subject`` for "Subject Heading",
    ``publish_date`` for "Citation Publish Date", ``pages`` for "Page Number").
  * ``Comments Due Date`` is GONE. The FAA export has it for 427 of 431 rows;
    this file drops the column entirely, so the comment deadline now comes from
    the Federal Register document itself (``comments_close_on``), which is the
    authoritative source for it anyway.
  * ``Link for Submitting Public Comment`` is gone too, replaced by
    ``federal_register_link``.
  * Four derived columns are new: ``product_subtype`` (370/431),
    ``supersedes_ad`` (62/431, "extracted from the summary text" per the sheet's
    own banner), ``mentions_serial_numbers`` (yes/no, 16 yes) and
    ``federal_register_link``.
  * ``ad_number`` is filled for only 21 rows here against 431 in the FAA export:
    the export carries the *proposed* AD number, this file carries the final one
    where a final rule exists. They are different facts and are not merged.

THE LINK COLUMN IS NOT A LINK
  The measured fill rate of ``federal_register_link`` is 406/431, but only
  **237** of those cells contain a Federal Register URL at all. 169 cells hold
  pasted document or withdrawal prose instead, and 15 of the 237 hold two URLs
  (the notice plus its later withdrawal). Three permalink shapes appear --
  ``/documents/YYYY/MM/DD/<number>/<slug>``, ``/d/<number>`` and
  ``/a/<number>`` -- all of which are official. So the column is treated as a
  *candidate* source of document numbers, parsed from the URL path, and every
  candidate is confirmed against the Federal Register index before use. The
  resolution itself lives in :mod:`engine.ingest.nprm_text`.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from engine.ingest.docket_ids import normalize_sheet_docket
from engine.ingest.make_normalize import normalize_make
from engine.ingest.nprm_xlsx import (
    NPRMRow,
    SpreadsheetError,
    clean_cell,
    page_range,
    parse_date,
)

#: ISO text, verified on 431/431 rows of both date columns.
DATE_FORMAT = "%Y-%m-%d"
SHEET_NAME = "nprm_index"
#: Row 0 is a banner sentence, row 1 is blank, row 2 is the header.
HEADER_ROW_INDEX = 2

_REQUIRED_COLUMNS = (
    "record_id",
    "docket_number",
    "subject",
    "issue_date",
    "publish_date",
    "make",
    "model",
    "product_type",
    "product_subtype",
    "ad_number",
    "supersedes_ad",
    "mentions_serial_numbers",
    "citation",
    "pages",
    "federal_register_link",
    "summary",
)

#: Federal Register permalink shapes, matched on the URL *path* rather than by
#: pattern-matching the surrounding prose. All three are official forms the
#: site itself issues, and all three appear in this column.
_FR_DOCUMENT_URL = re.compile(
    r"federalregister\.gov/"
    r"(?:documents/\d{4}/\d{2}/\d{2}/([^/\s\"'\\]+)/"  # long permalink
    r"|(?:d|a)/([^/\s\"'\\?#]+))"  # short permalinks
)
#: A page cell holds one or two integers. The wrapper varies ("(Pages 1-2)",
#: "Pages 1-2", "(Page 1)", and one cell with a doubled closing paren), so the
#: integers are read positionally instead of the wrapper being matched.
_PAGE_NUMBER = re.compile(r"\d+")

#: The sheet spells this column yes/no. Anything else is reported, not guessed.
_YES = "yes"
_NO = "no"


@dataclass
class IndexRow(NPRMRow):
    """One row of ``nprm_index``: an :class:`NPRMRow` plus the index's own keys.

    ``record_id`` is the organizers' stable key and is used as the graph's
    ``nprm_id`` (see :attr:`nprm_id`).
    """

    record_id: str = ""
    product_subtype: str | None = None
    supersedes_ad: str | None = None
    mentions_serial_numbers: bool | None = None
    federal_register_link: str | None = None
    #: Document numbers parsed out of ``federal_register_link``, in cell order.
    #: Candidates only: confirmed against the Federal Register index elsewhere.
    link_document_numbers: list[str] = field(default_factory=list)
    pages_raw: str | None = None
    start_page: int | None = None
    end_page: int | None = None

    @property
    def nprm_id(self) -> str:
        """The graph key for this proposal.

        ``record_id`` and not docket+date, for two reasons. It is the key the
        organizers publish and ask citations to name, and it contains no ``@``
        -- a citation id is ``{paragraph_id}@{char_start}-{char_end}`` and is
        split on its first ``@``, so an id carrying one would silently truncate
        every citation built on this document.
        """
        return self.record_id


@dataclass
class IndexLoad:
    rows: list[IndexRow]
    date_parse_failures: list[str] = field(default_factory=list)
    rows_without_docket: list[str] = field(default_factory=list)
    rows_with_multiple_dockets: list[str] = field(default_factory=list)
    #: record_ids whose link cell is empty.
    rows_without_link: list[str] = field(default_factory=list)
    #: record_ids whose link cell is non-empty but holds no Federal Register URL.
    rows_with_link_but_no_document_number: list[str] = field(default_factory=list)
    #: record_ids whose link cell holds more than one document number.
    rows_with_several_document_numbers: list[str] = field(default_factory=list)
    rows_without_page_numbers: list[str] = field(default_factory=list)
    unexpected_serial_number_flags: list[str] = field(default_factory=list)

    def as_report(self) -> dict[str, object]:
        return {
            "rows": len(self.rows),
            "unique_dockets": len({d for r in self.rows for d in r.docket_ids}),
            "date_parse_failures": self.date_parse_failures,
            "rows_without_docket": self.rows_without_docket,
            "rows_with_multiple_dockets": self.rows_with_multiple_dockets,
            "rows_without_link": self.rows_without_link,
            "rows_with_link_but_no_document_number": (
                self.rows_with_link_but_no_document_number
            ),
            "rows_with_several_document_numbers": self.rows_with_several_document_numbers,
            "rows_without_page_numbers": self.rows_without_page_numbers,
            "unexpected_serial_number_flags": self.unexpected_serial_number_flags,
            "mentions_serial_numbers_yes": sum(
                1 for r in self.rows if r.mentions_serial_numbers
            ),
            "with_supersedes_ad": sum(1 for r in self.rows if r.supersedes_ad),
            "with_ad_number": sum(1 for r in self.rows if r.ad_number),
            "with_model": sum(1 for r in self.rows if r.model_text),
        }


def federal_register_document_numbers(cell: object) -> list[str]:
    """Every Federal Register document number in one link cell, in order.

    Parsed from the URL path, which is structured, and de-duplicated. Returns
    ``[]`` for a cell with no Federal Register URL -- 169 of these cells hold
    pasted document text rather than a link, and inventing a number from that
    prose is exactly what this must not do.
    """
    if cell in (None, ""):
        return []
    found = (a or b for a, b in _FR_DOCUMENT_URL.findall(str(cell)))
    return list(dict.fromkeys(found))


def page_numbers(cell: object) -> tuple[int | None, int | None]:
    """``(start, end)`` Federal Register pages from a ``pages`` cell.

    The wrapper is inconsistent across the file, so the integers are read in
    order: one integer means a single page and ``end == start``.
    """
    if cell in (None, ""):
        return None, None
    numbers = [int(n) for n in _PAGE_NUMBER.findall(str(cell))]
    if not numbers:
        return None, None
    if len(numbers) == 1:
        return numbers[0], numbers[0]
    return numbers[0], numbers[1]


def _serial_flag(raw: object, record_id: str, unexpected: list[str]) -> bool | None:
    text = (clean_cell(raw) or "").lower()
    if text == _YES:
        return True
    if text == _NO:
        return False
    unexpected.append(f"{record_id}: mentions_serial_numbers={raw!r}")
    return None


def load_challenge_index(path: Path) -> IndexLoad:
    """Parse every data row of ``nprm_index``. Nothing is skipped."""
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        workbook.close()
        raise SpreadsheetError(
            f"{path} has no {SHEET_NAME!r} sheet; it has {workbook.sheetnames}"
        )
    raw_rows = list(workbook[SHEET_NAME].iter_rows(values_only=True))
    workbook.close()
    if len(raw_rows) <= HEADER_ROW_INDEX:
        raise SpreadsheetError(f"{path}:{SHEET_NAME} has no header row")

    header = [clean_cell(h) for h in raw_rows[HEADER_ROW_INDEX]]
    missing = [c for c in _REQUIRED_COLUMNS if c not in header]
    if missing:
        raise SpreadsheetError(
            f"{path}:{SHEET_NAME} is missing expected columns {missing}"
        )
    index = {name: position for position, name in enumerate(header) if name}

    load = IndexLoad(rows=[])
    seen: set[str] = set()
    for offset, raw in enumerate(raw_rows[HEADER_ROW_INDEX + 1 :], start=HEADER_ROW_INDEX + 2):
        if not any(cell not in (None, "") for cell in raw):
            continue

        def cell(name: str, raw=raw) -> object:
            return raw[index[name]]

        record_id = clean_cell(cell("record_id"))
        if not record_id:
            raise SpreadsheetError(f"{path}:{SHEET_NAME} row {offset} has no record_id")
        if record_id in seen:
            raise SpreadsheetError(
                f"{path}:{SHEET_NAME} row {offset} repeats record_id {record_id}"
            )
        seen.add(record_id)

        dockets, docket_raw = normalize_sheet_docket(str(cell("docket_number") or ""))
        if not dockets:
            load.rows_without_docket.append(record_id)
        if len(dockets) > 1:
            load.rows_with_multiple_dockets.append(record_id)

        link = clean_cell(cell("federal_register_link"))
        numbers = federal_register_document_numbers(cell("federal_register_link"))
        if not link:
            load.rows_without_link.append(record_id)
        elif not numbers:
            load.rows_with_link_but_no_document_number.append(record_id)
        if len(numbers) > 1:
            load.rows_with_several_document_numbers.append(record_id)

        start, end = page_numbers(cell("pages"))
        if start is None:
            load.rows_without_page_numbers.append(record_id)

        make_raw = clean_cell(cell("make"))
        load.rows.append(
            IndexRow(
                row_number=offset,
                docket_ids=dockets,
                docket_raw=docket_raw,
                subject_heading=clean_cell(cell("subject")),
                issue_date=parse_date(
                    cell("issue_date"),
                    f"{record_id} issue_date",
                    load.date_parse_failures,
                    (DATE_FORMAT,),
                ),
                ad_number=clean_cell(cell("ad_number")),
                publish_date=parse_date(
                    cell("publish_date"),
                    f"{record_id} publish_date",
                    load.date_parse_failures,
                    (DATE_FORMAT,),
                ),
                # This sheet has no comments-due column; the deadline comes from
                # the Federal Register document instead.
                comments_due_date=None,
                make=normalize_make(make_raw),
                make_raw=make_raw,
                model_text=clean_cell(cell("model")),
                product_type=clean_cell(cell("product_type")),
                comment_url=None,
                citation=clean_cell(cell("citation")),
                page_range=page_range(cell("pages")),
                summary=clean_cell(cell("summary")),
                record_id=record_id,
                product_subtype=clean_cell(cell("product_subtype")),
                supersedes_ad=clean_cell(cell("supersedes_ad")),
                mentions_serial_numbers=_serial_flag(
                    cell("mentions_serial_numbers"),
                    record_id,
                    load.unexpected_serial_number_flags,
                ),
                federal_register_link=link,
                link_document_numbers=numbers,
                pages_raw=clean_cell(cell("pages")),
                start_page=start,
                end_page=end,
            )
        )
    return load


def total_pages(rows: list[IndexRow]) -> dict[str, object]:
    """Federal Register page count over these rows, from the ``pages`` column.

    ``end - start + 1`` per row, inclusive, because a notice printed on pages
    66019-66022 occupies four pages. A row whose pages did not parse is counted
    in ``rows_without_pages`` and contributes nothing, so the total is a floor.

    Federal Register pages are shared: two notices can begin and end on the same
    physical page, so the sum double counts the shared ones. ``distinct_pages``
    is the honest number -- the size of the union of every (volume-free) page
    number touched -- and both are reported rather than one being presented as
    the other.
    """
    summed = 0
    touched: set[tuple[int, int]] = set()
    without = 0
    for row in rows:
        if row.start_page is None or row.end_page is None:
            without += 1
            continue
        start, end = min(row.start_page, row.end_page), max(row.start_page, row.end_page)
        summed += end - start + 1
        year = row.publish_date.year if row.publish_date else 0
        touched.update((year, page) for page in range(start, end + 1))
    return {
        "rows": len(rows),
        "rows_without_pages": without,
        "summed_pages": summed,
        "distinct_pages": len(touched),
    }


def _publish_years(rows: list[IndexRow]) -> list[int]:
    return sorted({row.publish_date.year for row in rows if row.publish_date})


def publish_year_range(rows: list[IndexRow]) -> tuple[int, int] | None:
    years = _publish_years(rows)
    if not years:
        return None
    return years[0], years[-1]
